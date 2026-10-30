"""ShadeInstall - Interní Aplikace."""

import re
import sys
from copy import copy
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tkinter import BooleanVar, filedialog, messagebox
import webbrowser

import customtkinter as ctk
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from PIL import Image


def resource_path(filename: str) -> Path:
    """Najde soubor vedle skriptu i uvnitř PyInstaller EXE."""
    base_path = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base_path / filename


def parse_dimension(value: str, name: str) -> Decimal:
    try:
        dimension = Decimal(value.strip().replace(",", "."))
    except InvalidOperation as error:
        raise ValueError(f"Zadejte platnou číselnou hodnotu pro {name}.") from error
    if not dimension.is_finite() or dimension <= 0:
        raise ValueError(f"{name.capitalize()} musí být kladné číslo.")
    return dimension


def parse_count(value: str) -> int:
    try:
        count = int(value.strip())
    except ValueError as error:
        raise ValueError("Počet musí být kladné celé číslo.") from error
    if count <= 0:
        raise ValueError("Počet musí být kladné celé číslo.")
    return count


def parse_discount(value: str) -> Decimal:
    value = value.strip()
    if not value:
        return Decimal("0")
    try:
        percent = Decimal(value.replace(",", "."))
    except InvalidOperation as error:
        raise ValueError("Sleva musí být číslo od 0 do 100 %.") from error
    if not percent.is_finite() or not 0 <= percent <= 100:
        raise ValueError("Sleva musí být číslo od 0 do 100 %.")
    return percent / 100


def parse_transport_distance(value: str) -> Decimal:
    value = value.strip()
    if not value:
        return Decimal("0")
    try:
        distance = Decimal(value.replace(",", "."))
    except InvalidOperation as error:
        raise ValueError("Vzdálenost musí být nezáporné číslo v kilometrech.") from error
    if not distance.is_finite() or distance < 0:
        raise ValueError("Vzdálenost musí být nezáporné číslo v kilometrech.")
    return distance


def calculate_holder_count(width_cm: Decimal) -> int:
    if width_cm <= 160:
        return 2
    if width_cm <= 240:
        return 3
    if width_cm <= 320:
        return 4
    if width_cm <= 420:
        return 5
    return 6


def apply_margin(price: int, enabled: bool) -> int:
    if not enabled:
        return price
    return (price * 11 + 9) // 10


def _shift_formula_rows(
    formula: str, delete_start: int, delete_end: int, deleted_count: int
) -> str:
    reference_pattern = re.compile(
        r"(?<![A-Z0-9_.])(?:(?P<range>\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+)"
        r"|(?P<cell>\$?[A-Z]{1,3}\$?\d+))(?![A-Z0-9_])",
        re.IGNORECASE,
    )
    cell_pattern = re.compile(r"(\$?[A-Z]{1,3}\$?)(\d+)", re.IGNORECASE)

    def shift_cell(reference: str) -> str:
        match = cell_pattern.fullmatch(reference)
        if match is None:
            return reference
        row = int(match.group(2))
        if delete_start <= row <= delete_end:
            return "#REF!"
        if row > delete_end:
            row -= deleted_count
        return f"{match.group(1)}{row}"

    def replace_reference(match: re.Match) -> str:
        cell_range = match.group("range")
        if cell_range is None:
            return shift_cell(match.group("cell"))

        first_ref, last_ref = cell_range.split(":")
        first_match = cell_pattern.fullmatch(first_ref)
        last_match = cell_pattern.fullmatch(last_ref)
        first_row = int(first_match.group(2))
        last_row = int(last_match.group(2))

        if first_row < delete_start <= last_row <= delete_end:
            last_row = delete_start - 1
            last_ref = f"{last_match.group(1)}{last_row}"
        elif delete_start <= first_row <= delete_end < last_row:
            first_ref = f"{first_match.group(1)}{delete_start}"
            last_ref = f"{last_match.group(1)}{last_row - deleted_count}"
        elif delete_start <= first_row <= last_row <= delete_end:
            return "#REF!"
        else:
            first_ref = shift_cell(first_ref)
            last_ref = shift_cell(last_ref)
        return f"{first_ref}:{last_ref}"

    return reference_pattern.sub(replace_reference, formula)


def _delete_unused_item_rows(worksheet, item_count: int) -> None:
    deleted_count = MAX_BLIND_ITEMS - item_count
    if deleted_count <= 0:
        return

    delete_start = ITEM_START_ROW + item_count
    delete_end = ITEM_START_ROW + MAX_BLIND_ITEMS - 1
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.data_type == "f" and isinstance(cell.value, str):
                cell.value = _shift_formula_rows(
                    cell.value, delete_start, delete_end, deleted_count
                )

    merges_to_move = []
    for merged_range in list(worksheet.merged_cells.ranges):
        min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
        if min_row <= delete_end and max_row >= delete_start:
            raise ValueError("Řádky položek zasahují do sloučených buněk šablony.")
        if min_row > delete_end:
            merges_to_move.append((min_col, min_row, max_col, max_row))
            worksheet.unmerge_cells(str(merged_range))

    row_dimensions = {
        index: copy(dimension)
        for index, dimension in worksheet.row_dimensions.items()
    }
    print_area = worksheet.print_area
    images_to_shift = []
    for image in worksheet._images:
        anchor = image.anchor
        start_marker = anchor._from
        end_marker = getattr(anchor, "to", None)
        if start_marker.row >= delete_start and (
            end_marker is None or end_marker.row >= delete_start
        ):
            images_to_shift.append((start_marker, end_marker))
        elif end_marker is not None and end_marker.row >= delete_start:
            raise ValueError(
                "Obrázek překrývá oblast mazaných řádků šablony."
            )

    worksheet.delete_rows(delete_start, deleted_count)

    for start_marker, end_marker in images_to_shift:
        start_marker.row -= deleted_count
        if end_marker is not None:
            end_marker.row -= deleted_count

    for index in list(worksheet.row_dimensions):
        del worksheet.row_dimensions[index]
    for index, dimension in row_dimensions.items():
        if delete_start <= index <= delete_end:
            continue
        new_index = index - deleted_count if index > delete_end else index
        dimension.index = new_index
        worksheet.row_dimensions[new_index] = dimension

    for min_col, min_row, max_col, max_row in merges_to_move:
        worksheet.merge_cells(
            start_row=min_row - deleted_count,
            start_column=min_col,
            end_row=max_row - deleted_count,
            end_column=max_col,
        )

    if print_area:
        area_reference = str(print_area).rsplit("!", 1)[-1]
        min_col, min_row, max_col, max_row = range_boundaries(area_reference)
        if min_row > delete_end:
            min_row -= deleted_count
        if max_row > delete_end:
            max_row -= deleted_count
        worksheet.print_area = (
            f"'{worksheet.title}'!${get_column_letter(min_col)}${min_row}:"
            f"${get_column_letter(max_col)}${max_row}"
        )


TEMPLATE_FILES = {
    "Zetta 90 - Standart": "data/vzor_cenik_zetta90.xlsx",
    "Zetta 90 - RAL": "data/vzor_cenik_Zetta90_ral.xlsx",
}
PRICE_FILES = {
    "Zetta 90 - Standart": "data/ceny_zetta90.xlsx",
    "Zetta 90 - RAL": "data/ceny_zetta90.xlsx",
}
ITEM_START_ROW = 23
MAX_BLIND_ITEMS = 14
MIN_ALLOWED_WIDTH_CM = Decimal("60")
MAX_ALLOWED_HEIGHT_CM = Decimal("400")
ROUTE_PLANNER_URL = (
    "https://mapy.com/cs/turisticka?planovani-trasy&rc=9hFPgxYGxI&rs=addr"
    "&rs=&ri=8991762&ri=&mrp=%7B%22c%22%3A111%2C%22dt%22%3A%22%22%2C%22d%22"
    "%3Atrue%7D&xc=%5B%5D&rut=1&rbf=alf&x=15.6252330&y=49.8022514&z=8"
)


def open_route_planner() -> None:
    webbrowser.open(ROUTE_PLANNER_URL, new=2)


def lookup_price(blind_type: str, width_mm: Decimal, height_mm: Decimal) -> int:
    price_file = PRICE_FILES.get(blind_type)
    if price_file is None:
        raise ValueError(f"Pro žaluzii {blind_type} není k dispozici ceník.")

    price_path = resource_path(price_file)
    if not price_path.exists():
        raise FileNotFoundError(f"Chybí cenová tabulka {price_file}.")

    workbook = load_workbook(price_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        widths = [worksheet.cell(row=1, column=column).value
                  for column in range(2, worksheet.max_column + 1)]
        heights = [worksheet.cell(row=row, column=1).value
                   for row in range(2, worksheet.max_row + 1)]
        selected_width = next(
            (value for value in widths if Decimal(str(value)) >= width_mm), None
        )
        selected_height = next(
            (value for value in heights if Decimal(str(value)) >= height_mm), None
        )
        if selected_width is None or selected_height is None:
            raise ValueError("Zadané rozměry přesahují rozsah cenové tabulky.")

        column = widths.index(selected_width) + 2
        row = heights.index(selected_height) + 2
        price = worksheet.cell(row=row, column=column).value
        if not isinstance(price, (int, float)):
            raise ValueError("Pro zadané rozměry není v cenové tabulce cena.")
        return int(price)
    finally:
        workbook.close()


def create_pricelist(
    output_path: Path,
    items: list[dict],
    quote_number: str,
    customer_name: str,
    installation_location: str,
    blind_type: str,
    blind_discount: Decimal,
    accessory_discount: Decimal,
    transport_cost: Decimal,
) -> None:
    """Vyplní kopii šablony zadanými parametry a uloží ji jako nový soubor."""
    if not 1 <= len(items) <= MAX_BLIND_ITEMS:
        raise ValueError(f"Ceník musí obsahovat 1 až {MAX_BLIND_ITEMS} žaluzií.")
    template_file = TEMPLATE_FILES.get(blind_type)
    if template_file is None:
        raise ValueError(f"Pro žaluzii {blind_type} není k dispozici vzor.")
    template_path = resource_path(template_file)
    if not template_path.exists():
        raise FileNotFoundError(f"Chybí šablona {template_file}.")
    workbook = load_workbook(template_path)
    try:
        worksheet = workbook.active
        worksheet["K9"] = quote_number
        worksheet["E9"] = customer_name
        worksheet["E10"] = installation_location
        worksheet["J22"] = float(blind_discount)
        worksheet["J42"] = float(accessory_discount)
        transport_row = next(
            (
                row
                for row in range(1, worksheet.max_row + 1)
                if isinstance(worksheet.cell(row=row, column=2).value, str)
                and worksheet.cell(row=row, column=2).value.strip().casefold()
                .startswith("doprava")
            ),
            None,
        )
        if transport_row is None:
            raise ValueError("Vzor ceníku neobsahuje řádek Doprava.")
        worksheet.cell(row=transport_row, column=9, value=float(transport_cost))
        for item_index, item in enumerate(items):
            row = ITEM_START_ROW + item_index
            worksheet[f"B{row}"] = item["holder_count"]
            worksheet[f"E{row}"] = float(item["width_cm"])
            worksheet[f"F{row}"] = float(item["height_cm"])
            worksheet[f"G{row}"] = item["count"]
            worksheet[f"I{row}"] = item["price"]
        _delete_unused_item_rows(worksheet, len(items))
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
        workbook.save(output_path)
    finally:
        workbook.close()


BLIND_TYPES = ["Zetta 90 - Standart", "Zetta 90 - RAL"]


class ShadeInstallApp(ctk.CTk):
    def __init__(self) -> None:
        super().__init__()
        self.title("ShadeInstall - Interní Aplikace")
        self.geometry("1365x910")
        self.minsize(850, 600)
        self.configure(fg_color="#f8f9fa")
        self.iconbitmap(str(resource_path("logo.ico")))
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self.sidebar_frame = ctk.CTkFrame(
            self, width=260, corner_radius=0, fg_color="#ffffff"
        )
        self.sidebar_frame.grid(row=0, column=0, sticky="nsew")
        self.sidebar_frame.grid_propagate(False)
        self.sidebar_frame.grid_columnconfigure(0, weight=1)

        brand_frame = ctk.CTkFrame(self.sidebar_frame, fg_color="transparent")
        brand_frame.grid(row=0, column=0, padx=20, pady=(25, 20), sticky="w")
        self.brand_image = ctk.CTkImage(
            light_image=Image.open(resource_path("logo_mark.png")),
            size=(42, 42),
        )
        ctk.CTkLabel(brand_frame, text="", image=self.brand_image).pack(side="left")
        ctk.CTkLabel(
            brand_frame,
            text="ShadeInstall.cz",
            font=("Segoe UI", 16, "bold"),
            text_color="#0d6efd",
        ).pack(side="left", padx=(10, 0))
        ctk.CTkLabel(
            self.sidebar_frame,
            text="MODULY",
            font=("Segoe UI", 10, "bold"),
            text_color="#94a3b8",
        ).grid(row=1, column=0, padx=20, pady=(10, 5), sticky="w")

        ctk.CTkButton(
            self.sidebar_frame,
            text="Příprava cenové nabídky",
            fg_color="#0d6efd",
            hover_color="#0b5ed7",
            text_color="#ffffff",
            corner_radius=8,
            anchor="w",
            font=("Segoe UI", 13, "bold"),
            height=42,
        ).grid(row=2, column=0, padx=12, pady=4, sticky="ew")

        self.main_container = ctk.CTkFrame(
            self, fg_color="#f8f9fa", corner_radius=0
        )
        self.main_container.grid(row=0, column=1, sticky="nsew")
        self.main_container.grid_columnconfigure(0, weight=1)
        self.main_container.grid_rowconfigure(1, weight=1)

        self.top_bar = ctk.CTkFrame(
            self.main_container, height=65, fg_color="#ffffff", corner_radius=0
        )
        self.top_bar.grid(row=0, column=0, sticky="ew")
        self.top_bar.grid_propagate(False)
        self.top_bar.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            self.top_bar,
            text="Příprava cenové nabídky",
            font=("Segoe UI", 18, "bold"),
            text_color="#1e293b",
        ).grid(row=0, column=0, padx=25, pady=15, sticky="w")

        content = ctk.CTkScrollableFrame(
            self.main_container,
            fg_color="#ffffff",
            corner_radius=12,
            scrollbar_button_color="#cbd5e1",
            scrollbar_button_hover_color="#94a3b8",
        )
        content.grid(row=1, column=0, padx=25, pady=25, sticky="nsew")

        customer_quote_frame = ctk.CTkFrame(content, fg_color="transparent")
        customer_quote_frame.pack(fill="x", padx=30, pady=(5, 2))
        customer_fields = []
        for column in range(3):
            customer_quote_frame.grid_columnconfigure(
                column, weight=1, uniform="customer-fields"
            )
            field_frame = ctk.CTkFrame(
                customer_quote_frame, fg_color="transparent"
            )
            field_frame.grid(
                row=0,
                column=column,
                padx=(0 if column == 0 else 6, 0 if column == 2 else 6),
                sticky="ew",
            )
            customer_fields.append(field_frame)
        self.entry_customer_name = self._add_text_field(
            customer_fields[0], "Jméno zákazníka:", width=130, padx=4
        )
        self.entry_installation_location = self._add_text_field(
            customer_fields[1], "Místo realizace:", width=130, padx=4
        )
        self.entry_quote_number = self._add_text_field(
            customer_fields[2], "Číslo cenové nabídky:", width=130, padx=4
        )
        ctk.CTkLabel(
            content,
            text="Výběr žaluzie:",
            font=("Segoe UI", 13),
            text_color="#475569",
        ).pack(anchor="w", padx=30, pady=(5, 2))
        self.blind_type = ctk.CTkOptionMenu(
            content,
            values=BLIND_TYPES,
            width=320,
            height=38,
            corner_radius=8,
            fg_color="#f8f9fa",
            button_color="#cbd5e1",
            button_hover_color="#94a3b8",
            text_color="#1e293b",
        )
        self.blind_type.pack(anchor="w", padx=30, pady=(0, 12))

        self.blind_items_frame = ctk.CTkFrame(content, fg_color="transparent")
        self.blind_items_frame.pack(fill="x", padx=30, pady=(8, 8))
        for column, label in enumerate(
            ("Žaluzie", "Šířka (cm)", "Výška (cm)", "Počet")
        ):
            ctk.CTkLabel(
                self.blind_items_frame,
                text=label,
                font=("Segoe UI", 12, "bold"),
                text_color="#475569",
            ).grid(row=0, column=column, padx=5, pady=(0, 4), sticky="w")
        self.blind_items_frame.grid_columnconfigure(0, minsize=90)
        for column in range(1, 4):
            self.blind_items_frame.grid_columnconfigure(column, weight=1)

        self.blind_items = []
        self.add_blind_button = ctk.CTkButton(
            content,
            text="Přidat další žaluzii",
            fg_color="#16803c",
            hover_color="#126c32",
            text_color="#ffffff",
            corner_radius=8,
            font=("Segoe UI", 12, "bold"),
            width=260,
            height=34,
            command=lambda: self.add_blind_item(focus_new=True),
        )
        self.add_blind_button.pack(anchor="w", padx=30, pady=(0, 12))
        transport_field = ctk.CTkFrame(content, fg_color="transparent")
        transport_field.pack(anchor="w", padx=30, pady=(0, 4))
        self.entry_transport_distance = self._add_text_field(
            transport_field,
            "Cena dopravy (km jedné trasy):",
            width=170,
            padx=4,
            entry_bottom_padding=2,
        )
        self.route_planner_link = ctk.CTkLabel(
            transport_field,
            text="Spočítat trasu na Mapy.cz",
            font=("Segoe UI", 11, "underline"),
            text_color="#0d6efd",
            fg_color="transparent",
            cursor="hand2",
        )
        self.route_planner_link.pack(anchor="w", padx=4, pady=0)
        self.route_planner_link.bind(
            "<Button-1>", lambda _event: open_route_planner()
        )
        discount_frame = ctk.CTkFrame(content, fg_color="transparent")
        discount_frame.pack(anchor="w", padx=30, pady=(0, 8))
        discount_frame.grid_columnconfigure(0, weight=1, uniform="discount-fields")
        discount_frame.grid_columnconfigure(1, weight=1, uniform="discount-fields")
        blind_discount_field = ctk.CTkFrame(
            discount_frame, fg_color="transparent"
        )
        blind_discount_field.grid(row=0, column=0, padx=(0, 3), sticky="ew")
        accessory_discount_field = ctk.CTkFrame(
            discount_frame, fg_color="transparent"
        )
        accessory_discount_field.grid(row=0, column=1, padx=(3, 0), sticky="ew")
        self.entry_blind_discount = self._add_text_field(
            blind_discount_field, "Sleva na žaluzie (%):", width=130, padx=4
        )
        self.entry_accessory_discount = self._add_text_field(
            accessory_discount_field,
            "Sleva příslušenství (%):",
            width=130,
            padx=4,
        )
        self.margin_enabled = BooleanVar(master=self, value=True)
        self.margin_checkbox = ctk.CTkCheckBox(
            content,
            text="Započítat +10% nad ceník",
            variable=self.margin_enabled,
            onvalue=True,
            offvalue=False,
            checkbox_width=18,
            checkbox_height=18,
            fg_color="#16803c",
            hover_color="#126c32",
            text_color="#1e293b",
            font=("Segoe UI", 13),
        )
        self.margin_checkbox.pack(anchor="w", padx=30, pady=(0, 12))
        self.add_blind_item(focus_new=False)
        self.price_result = ctk.CTkLabel(
            content,
            text="",
            font=("Segoe UI", 16, "bold"),
            text_color="#16803c",
        )
        self.price_result.pack(anchor="w", padx=30, pady=(0, 8))
        self.create_button = ctk.CTkButton(
            content,
            text="Vytvořit cenovou nabídku",
            fg_color="#0d6efd",
            hover_color="#0b5ed7",
            text_color="#ffffff",
            corner_radius=8,
            font=("Segoe UI", 13, "bold"),
            width=320,
            height=42,
            command=self.calculate_and_write,
        )
        self.create_button.pack(anchor="w", padx=30, pady=10)

        self._bind_enter(self.entry_customer_name, self.entry_installation_location)
        self._bind_enter(self.entry_installation_location, self.entry_quote_number)
        self._bind_enter(self.entry_quote_number, self.blind_type)
        self._bind_enter(self.blind_type, self.blind_items[0]["width"])
        self.add_blind_button.bind(
            "<Return>", self._handle_add_blind_enter, add="+"
        )
        self.create_button.bind(
            "<Return>", self._handle_create_enter, add="+"
        )

    @staticmethod
    def _add_text_field(
        parent: ctk.CTkFrame,
        label: str,
        width: int = 320,
        padx: int = 30,
        entry_bottom_padding: int = 12,
    ) -> ctk.CTkEntry:
        ctk.CTkLabel(
            parent, text=label, font=("Segoe UI", 13), text_color="#475569"
        ).pack(anchor="w", padx=padx, pady=(5, 2))
        entry = ctk.CTkEntry(
            parent,
            width=width,
            height=38,
            corner_radius=8,
            fg_color="#f8f9fa",
            border_color="#cbd5e1",
            text_color="#1e293b",
        )
        entry.pack(
            fill="x",
            expand=True,
            padx=padx,
            pady=(0, entry_bottom_padding),
        )
        return entry

    @staticmethod
    def _create_item_entry(parent: ctk.CTkFrame, width: int) -> ctk.CTkEntry:
        return ctk.CTkEntry(
            parent,
            width=width,
            height=36,
            corner_radius=8,
            fg_color="#f8f9fa",
            border_color="#cbd5e1",
            text_color="#1e293b",
        )

    @staticmethod
    def _move_focus(_event, target) -> str:
        target.focus_set()
        return "break"

    def _bind_enter(self, widget, target) -> None:
        widget.bind(
            "<Return>",
            lambda event, next_widget=target: self._move_focus(
                event, next_widget
            ),
            add="+",
        )

    def _focus_after_item(self, event, item_index: int) -> str:
        next_index = item_index + 1
        if next_index < len(self.blind_items):
            target = self.blind_items[next_index]["width"]
        elif len(self.blind_items) < MAX_BLIND_ITEMS:
            target = self.add_blind_button
        else:
            target = self.create_button
        return self._move_focus(event, target)

    def _handle_add_blind_enter(self, _event) -> str:
        self.add_blind_item(focus_new=True)
        return "break"

    def _handle_create_enter(self, _event) -> str:
        self.calculate_and_write()
        return "break"

    @staticmethod
    def _update_width_warning(width_entry, warning_frame) -> None:
        try:
            width_cm = parse_dimension(width_entry.get(), "šířku")
        except ValueError:
            warning_frame.grid_remove()
            return

        if width_cm < MIN_ALLOWED_WIDTH_CM:
            warning_frame.grid()
        else:
            warning_frame.grid_remove()

    @staticmethod
    def _update_height_warning(height_entry, warning_frame) -> None:
        try:
            height_cm = parse_dimension(height_entry.get(), "výšku")
        except ValueError:
            warning_frame.grid_remove()
            return

        if height_cm > MAX_ALLOWED_HEIGHT_CM:
            warning_frame.grid()
        else:
            warning_frame.grid_remove()

    def add_blind_item(self, focus_new: bool = True) -> None:
        if len(self.blind_items) >= MAX_BLIND_ITEMS:
            return

        item_number = len(self.blind_items) + 1
        row_frame = ctk.CTkFrame(self.blind_items_frame, fg_color="transparent")
        row_frame.grid(
            row=item_number,
            column=0,
            columnspan=4,
            padx=0,
            pady=3,
            sticky="ew",
        )
        row_frame.grid_columnconfigure(0, minsize=90)
        for column in range(1, 4):
            row_frame.grid_columnconfigure(column, weight=1)

        ctk.CTkLabel(
            row_frame,
            text=str(item_number),
            font=("Segoe UI", 13),
            text_color="#475569",
        ).grid(row=0, column=0, padx=5, sticky="nw")
        width_field = ctk.CTkFrame(row_frame, fg_color="transparent")
        width_field.grid(row=0, column=1, padx=5, sticky="nw")
        width_entry = self._create_item_entry(width_field, 170)
        width_entry.grid(row=0, column=0, padx=(0, 8), sticky="w")
        warning_frame = ctk.CTkFrame(width_field, fg_color="transparent")
        warning_frame.grid(row=1, column=0, pady=(2, 0), sticky="w")
        ctk.CTkLabel(
            warning_frame,
            text="!",
            font=("Segoe UI", 14, "bold"),
            text_color="#b42318",
        ).pack(side="left", padx=(0, 4))
        ctk.CTkLabel(
            warning_frame,
            text="Žaluzie je pod povolený rozměr",
            font=("Segoe UI", 10),
            text_color="#b42318",
            wraplength=150,
            justify="left",
        ).pack(side="left")
        warning_frame.grid_remove()
        update_warning = lambda _event: self._update_width_warning(
            width_entry, warning_frame
        )
        width_entry.bind("<KeyRelease>", update_warning, add="+")
        width_entry.bind("<FocusOut>", update_warning, add="+")
        height_field = ctk.CTkFrame(row_frame, fg_color="transparent")
        height_field.grid(row=0, column=2, padx=5, sticky="nw")
        height_entry = self._create_item_entry(height_field, 170)
        height_entry.grid(row=0, column=0, padx=(0, 8), sticky="w")
        height_warning = ctk.CTkFrame(height_field, fg_color="transparent")
        height_warning.grid(row=1, column=0, pady=(2, 0), sticky="w")
        ctk.CTkLabel(
            height_warning,
            text="!",
            font=("Segoe UI", 14, "bold"),
            text_color="#b42318",
        ).pack(side="left", padx=(0, 4))
        ctk.CTkLabel(
            height_warning,
            text="Žaluzie je nad povolený rozměr",
            font=("Segoe UI", 10),
            text_color="#b42318",
            wraplength=150,
            justify="left",
        ).pack(side="left")
        height_warning.grid_remove()
        update_height_warning = lambda _event: self._update_height_warning(
            height_entry, height_warning
        )
        height_entry.bind("<KeyRelease>", update_height_warning, add="+")
        height_entry.bind("<FocusOut>", update_height_warning, add="+")
        count_entry = self._create_item_entry(row_frame, 130)
        count_entry.grid(row=0, column=3, padx=5, sticky="nw")
        self._bind_enter(width_entry, height_entry)
        self._bind_enter(height_entry, count_entry)
        count_entry.bind(
            "<Return>",
            lambda event, index=item_number - 1: self._focus_after_item(
                event, index
            ),
            add="+",
        )
        self.blind_items.append(
            {
                "frame": row_frame,
                "width": width_entry,
                "width_warning": warning_frame,
                "height": height_entry,
                "height_warning": height_warning,
                "count": count_entry,
            }
        )
        if len(self.blind_items) >= MAX_BLIND_ITEMS:
            self.add_blind_button.configure(state="disabled")
        if focus_new:
            width_entry.focus_set()

    def calculate_and_write(self) -> None:
        customer_name = self.entry_customer_name.get().strip()
        if not customer_name:
            messagebox.showerror(
                "Chybí jméno zákazníka",
                "Zadejte jméno zákazníka.",
                parent=self,
            )
            return
        customer_surname = customer_name.split()[-1]
        customer_label = f"p. {customer_surname}"
        customer_name = re.sub(
            r'[<>:"/\\|?*\x00-\x1f]', "_", customer_name
        ).strip(" .")
        if not customer_name:
            messagebox.showerror(
                "Neplatné jméno zákazníka",
                "Zadejte jméno obsahující znaky použitelné v názvu souboru.",
                parent=self,
            )
            return

        quote_number = self.entry_quote_number.get().strip()
        if not quote_number:
            messagebox.showerror(
                "Chybí číslo cenové nabídky",
                "Zadejte číslo cenové nabídky.",
                parent=self,
            )
            return
        quote_number_for_filename = quote_number
        quote_number = f"{quote_number}-{datetime.now():%y}"
        installation_location = self.entry_installation_location.get().strip()
        try:
            blind_discount = parse_discount(self.entry_blind_discount.get())
            accessory_discount = parse_discount(
                self.entry_accessory_discount.get()
            )
            transport_distance = parse_transport_distance(
                self.entry_transport_distance.get()
            )
        except ValueError as error:
            messagebox.showerror("Neplatné údaje", str(error), parent=self)
            return
        transport_cost = transport_distance * 4 * 13

        items = []
        for item_number, fields in enumerate(self.blind_items, start=1):
            raw_values = (
                fields["width"].get().strip(),
                fields["height"].get().strip(),
                fields["count"].get().strip(),
            )
            if not any(raw_values):
                continue
            try:
                width_cm = parse_dimension(
                    raw_values[0], f"šířku žaluzie {item_number}"
                )
                height_cm = parse_dimension(
                    raw_values[1], f"výšku žaluzie {item_number}"
                )
                if height_cm > MAX_ALLOWED_HEIGHT_CM:
                    raise ValueError(
                        f"Výška žaluzie {item_number} nesmí překročit 400 cm."
                    )
                count = parse_count(raw_values[2])
            except ValueError as error:
                messagebox.showerror(
                    "Neplatné údaje", str(error), parent=self
                )
                return

            try:
                base_price = lookup_price(
                    self.blind_type.get(), width_cm * 10, height_cm * 10
                )
            except (FileNotFoundError, ValueError) as error:
                self.price_result.configure(text="")
                messagebox.showerror(
                    "Cenu se nepodařilo vypočítat",
                    f"Žaluzie {item_number}: {error}",
                    parent=self,
                )
                return
            price = apply_margin(base_price, self.margin_enabled.get())
            items.append(
                {
                    "width_cm": width_cm,
                    "height_cm": height_cm,
                    "holder_count": calculate_holder_count(width_cm),
                    "count": count,
                    "price": price,
                }
            )

        if not items:
            messagebox.showerror(
                "Chybí žaluzie",
                "Vyplňte parametry alespoň jedné žaluzie.",
                parent=self,
            )
            return

        total_price = sum(item["price"] * item["count"] for item in items)
        self.price_result.configure(
            text=f"Celková cena: {total_price:,} Kč".replace(",", " ")
        )
        default_filename = (
            f"{quote_number_for_filename}, p. {customer_surname}, "
            f"venkovní žaluzie {self.blind_type.get()}.xlsx"
        )

        selected_file = filedialog.asksaveasfilename(
            title="Uložit nový ceník",
            defaultextension=".xlsx",
            initialfile=default_filename,
            filetypes=(("Excel sešit", "*.xlsx"),),
            parent=self,
        )
        if not selected_file:
            return
        output_path = Path(selected_file)
        try:
            create_pricelist(
                output_path,
                items,
                quote_number,
                customer_label,
                installation_location,
                self.blind_type.get(),
                blind_discount,
                accessory_discount,
                transport_cost,
            )
        except Exception as error:
            messagebox.showerror(
                "Vytvoření ceníku se nezdařilo", str(error), parent=self
            )
            return

        messagebox.showinfo(
            "Ceník vytvořen", f"Nový ceník:\n{output_path}", parent=self
        )
        for entry in (
            self.entry_customer_name,
            self.entry_quote_number,
            self.entry_installation_location,
            self.entry_blind_discount,
            self.entry_accessory_discount,
            self.entry_transport_distance,
        ):
            entry.delete(0, "end")
        for item in self.blind_items[1:]:
            item["frame"].destroy()
        self.blind_items = self.blind_items[:1]
        for field in ("width", "height", "count"):
            self.blind_items[0][field].delete(0, "end")
        self.blind_items[0]["width_warning"].grid_remove()
        self.blind_items[0]["height_warning"].grid_remove()
        self.add_blind_button.configure(state="normal")
        self.price_result.configure(text="")
        self.blind_type.set(BLIND_TYPES[0])


def main() -> None:
    ctk.set_appearance_mode("light")
    ctk.set_default_color_theme("blue")
    ShadeInstallApp().mainloop()


if __name__ == "__main__":
    main()